"""Build CloudMLOps faculty presentation from the Shopez APSCHE template."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

REFERENCE = Path(
    r"c:\Users\chimm\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm"
    r"\LocalState\sessions\C9878F0B5382DCA9B1FA5A9F7323802AAF0D56B9\transfers\2026-39"
    r"\Shopez_APSCHE_Professional_Presentation (1).pptx"
)
OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "CloudMLOps_Faculty_Presentation.pptx"
FOOTER = "CloudMLOps  •  CIS 5690  •  "


def replace_in_slide(slide, mapping: dict[str, str]) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text
        for old, new in mapping.items():
            if old in text:
                shape.text = text.replace(old, new)
                text = shape.text


def set_text_containing(slide, needle: str, new_text: str) -> bool:
    for shape in slide.shapes:
        if shape.has_text_frame and needle in shape.text:
            shape.text = new_text
            return True
    return False


def global_footer_replace(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text
            if "APSCHE PROJECT" in t or "SHOPEZ" in t and "•" in t:
                # footer pattern: SHOPEZ  •  APSCHE PROJECT  •  N
                parts = t.split("•")
                if len(parts) >= 3 and parts[-1].strip().isdigit():
                    n = parts[-1].strip()
                    shape.text = f"{FOOTER}{n}"


def add_stack_textbox(slide) -> None:
    left = Inches(0.8)
    top = Inches(2.0)
    width = Inches(11.5)
    height = Inches(4.5)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = [
        ("Presentation tier", "React 19, Vite, React Router, Axios"),
        ("Business tier", "Python 3.12, FastAPI, SQLAlchemy 2, PyJWT, bcrypt"),
        ("Data tier", "PostgreSQL 16, Alembic migrations"),
        ("AI / ML", "Hugging Face Transformers, google/flan-t5-small, extractive fallback"),
        ("Document extraction", "PyMuPDF (PDF), python-docx (DOCX), plain text"),
        ("DevOps / Cloud", "Docker, GitHub Actions, Amazon ECR, AWS App Runner, RDS, S3, CloudWatch"),
        ("Infrastructure as code", "Terraform (VPC connector, App Runner, IAM)"),
        ("Quality", "pytest (171 tests), Ruff, CI Postgres integration job"),
    ]
    for i, (title, body) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{title}: {body}"
        p.level = 0
        p.font.size = Pt(14)


def main() -> int:
    if not REFERENCE.is_file():
        print(f"Reference PPT not found: {REFERENCE}", file=sys.stderr)
        return 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    prs = Presentation(str(OUTPUT))

    # Slide 1 — Title
    s1 = prs.slides[0]
    replace_in_slide(
        s1,
        {
            "SHOPEZ": "CloudMLOps",
            "E-Commerce Application": "AI Document Summarization Platform",
            "APSCHE • Team Project": "CIS 5690 • Cloud MLOps Project",
        },
    )
    set_text_containing(
        s1,
        "Team members",
        "Presented by:\nChimmili Nityanandh (GitHub: nitya200)\nchimmilinityanandh@gmail.com",
    )

    # Slide 2 — Overview
    s2 = prs.slides[1]
    replace_in_slide(
        s2,
        {
            "A simple, modern shopping experience": "Summarize long documents with AI and MLOps on AWS",
            "What is Shopez?": "What is CloudMLOps?",
            "An e-commerce application that helps users discover products, manage a cart and place orders.":
            "A three-tier web platform that turns PDF, DOCX, and TXT files (or pasted text) into "
            "searchable summaries stored in PostgreSQL, with admin metrics and model lifecycle controls.",
            "Project Approach": "Project Approach",
            "A team-developed application created to apply software development concepts in a practical project.":
            "End-to-end course project applying software architecture, AI integration, databases, "
            "automated testing, and cloud deployment (Docker → GitHub Actions → AWS App Runner).",
            "Target Experience": "Target Experience",
            "Easy navigation • Clear product presentation • Smooth shopping workflow":
            "Fast upload • Clear summaries • History & ratings • Secure admin dashboard • Live AWS URL",
            "SHOP SMART. LIVE BETTER.": "READ LESS. UNDERSTAND MORE.",
            "Fashion": "PDF",
            "Electronics": "DOCX",
            "Beauty": "TXT",
            "Home": "Paste",
            "Sports": "History",
            "Deals": "Admin",
        },
    )

    # Slide 3 — Problem & objectives
    s3 = prs.slides[2]
    set_text_containing(
        s3,
        "Customers need a convenient way",
        "Professionals and students receive long reports and research papers that take too much time "
        "to read in full. Manual skimming misses key points, and there is no single place to store, "
        "search, and rate summaries over time.",
    )
    set_text_containing(
        s3,
        "Traditional shopping methods",
        "Organizations need a secure, multi-user platform that extracts text from common document "
        "formats, generates abstractive or fast extractive summaries, persists results in a relational "
        "database, and deploys reliably on cloud infrastructure with CI/CD.",
    )
    set_text_containing(
        s3,
        "Build a user-friendly shopping platform",
        "• Build a three-tier summarization platform (React / FastAPI / PostgreSQL)\n"
        "• Support PDF, DOCX, TXT upload and raw text input\n"
        "• Integrate FLAN-T5-small with factory/strategy design patterns\n"
        "• Provide auth, history, feedback, and admin analytics\n"
        "• Deploy on AWS App Runner with Terraform and GitHub Actions",
    )

    # Slide 4 — Key features (8 tiles)
    s4 = prs.slides[3]
    replace_in_slide(
        s4,
        {
            "Core functionality of the Shopez application": "Core capabilities of the CloudMLOps platform",
            "User Access": "Authentication",
            "Registration, login and profile management": "Register, login, JWT sessions, role-based access",
            "Product Discovery": "Documents",
            "Browse categories, search and view details": "Upload PDF/DOCX/TXT, magic-byte validation, extraction",
            "Cart": "Summarization",
            "Add, remove and update cart items": "Short / medium / long summaries, download as .txt",
            "Orders": "History",
            "Checkout and order tracking details": "Search, paginate, and delete past summaries",
            "Reviews & Ratings": "Feedback",
            "Customers rate products and share reviews": "1–5 star ratings stored for quality metrics",
            "AI Chatbot": "AI engine",
            "Instant chat support for queries and help": "FLAN-T5 abstractive (local) + extractive on AWS for speed",
            "Admin (CRUD)": "Admin dashboard",
            "Full create, read, update and delete control": "Users, usage stats, ROUGE model approve/promote (UC-12/16)",
            "Responsive UI": "Cloud deployment",
            "Clean, adaptive layout across all devices": "Live on AWS App Runner + reference Netlify/Render demo",
        },
    )

    # Slide 5 — Modules
    s5 = prs.slides[4]
    replace_in_slide(
        s5,
        {
            "Organized around the main shopping workflow": "Organized around the summarization workflow",
            "USER": "CLIENT",
            "Login • Profile • Access": "React UI • Auth context • Protected routes",
            "PRODUCT": "API",
            "Catalogue • Search • Details": "FastAPI routers • Services • Repositories",
            "CART": "AI",
            "Add • Remove • Quantity": "Factory • Strategy • Map-reduce chunks",
            "ORDER": "DATA",
            "Checkout • Order details": "PostgreSQL • Alembic • S3/local storage",
            "ADMIN": "MLOPS",
            "Management • Updates": "CI/CD • ECR • Model registry • CloudWatch logs",
            "Browse  →  Select  →  Cart  →  Checkout  →  Order":
            "Upload  →  Extract  →  Summarize  →  Store  →  Rate / Download",
        },
    )

    # Slide 6 — Stack
    s6 = prs.slides[5]
    replace_in_slide(
        s6,
        {
            "The MERN-based stack powering Shopez": "Full stack from UI to AWS (see list below; diagram: MERN reference layout)",
        },
    )
    add_stack_textbox(s6)

    # Slide 7 — User journey
    s7 = prs.slides[6]
    replace_in_slide(
        s7,
        {
            "End-to-end shopping flow": "End-to-end summarization flow",
            "Open Shopez": "Open CloudMLOps",
            "Browse Products": "Upload / paste text",
            "View Details": "Choose length",
            "Add to Cart": "Generate summary",
            "Checkout": "View & download",
            "Place Order": "Rate & history",
        },
    )
    set_text_containing(
        s7,
        "Create Product",
        "Admin & MLOps flow\n\n"
        "Monitor ➜ Platform stats, usage metrics, rating distribution\n"
        "Manage users ➜ Activate/deactivate accounts, assign admin role\n"
        "Model lifecycle (UC-12) ➜ Review ROUGE scores, approve and promote a version\n"
        "Retraining job (UC-16) ➜ Trigger evaluation job, register new model version\n\n"
        "Bottom line:\nEvaluate → Approve → Promote → Production summarizer",
    )

    # Slide 8 — Testing
    s8 = prs.slides[7]
    replace_in_slide(
        s8,
        {
            "Focus on reliability of the main user journey": "Focus on quality, security, and deployment readiness",
            "Functional Testing": "Automated testing",
            "Login, product browsing, cart and order flow": "171 pytest cases: auth, documents, summarization, admin, migrations",
            "UI Testing": "CI pipeline",
            "Navigation, layout and responsive behaviour": "Ruff lint/format, Docker build, Postgres integration job",
            "Bug Fixing": "Live verification",
            "Issues identified during development were corrected": "AWS /health: database connected; App Runner URLs verified",
            "Result": "Result",
            "Core e-commerce workflow demonstrated successfully": "Three-tier app deployed on AWS with documented demo URLs",
        },
    )

    # Slide 9 — Future
    s9 = prs.slides[8]
    replace_in_slide(
        s9,
        {
            "Potential next steps for Shopez": "Potential next steps for CloudMLOps",
            "Online payment integration": "GPU App Runner or SageMaker for FLAN-T5 at scale",
            "Order tracking & notifications": "CloudWatch alarms and SNS alerts",
            "Wishlist & saved items": "Email verification on registration",
            "Personalized recommendations": "Fine-tune FLAN-T5 on rated summaries",
            "Advanced search & filters": "Full ROUGE vs DistilBART benchmark suite",
            "Analytics & admin dashboard": "OIDC-only deploy (remove access-key ECR fallback)",
        },
    )

    # Slide 10 — Conclusion
    s10 = prs.slides[9]
    set_text_containing(
        s10,
        "What We Achieved",
        "What We Achieved\n\n"
        "CloudMLOps is an AI document summarization platform with clean architecture, real PostgreSQL "
        "persistence, and production deployment on AWS App Runner.\n\n"
        "The project demonstrates Repository/Factory/Strategy patterns, FLAN-T5 integration, Alembic "
        "migrations, 171 automated tests, and a complete MLOps path from GitHub Actions to ECR.\n\n"
        "CloudMLOps shows how containerization, Terraform, and CI/CD deliver a faculty-reviewable live "
        "system—not only local code.",
    )
    set_text_containing(
        s10,
        "A strong foundation",
        "Live demo: https://asqmhsdwfs.us-east-2.awsapprunner.com (UI) • "
        "https://p3jivcdmbf.us-east-2.awsapprunner.com (API)",
    )
    set_text_containing(s10, "user-friendly e-commerce", "")

    # Slide 11 — Thank you
    s11 = prs.slides[10]
    replace_in_slide(s11, {"SHOPEZ": "CloudMLOps"})

    global_footer_replace(prs)

    # Second pass: any remaining SHOPEZ branding
    for slide in prs.slides:
        replace_in_slide(
            slide,
            {
                "SHOPEZ": "CloudMLOps",
                "Shopez": "CloudMLOps",
                "APSCHE PROJECT": "CIS 5690",
            },
        )

    prs.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size / 1024:.1f} KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
